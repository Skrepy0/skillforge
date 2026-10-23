"""Replay the real request through CLI, collection, generation, gates, and rendering."""

from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path

from pptx import Presentation
from skillforge.l3_orchestration.orchestrator import Orchestrator
from skillforge.l4_interface import cli, llm_client


def test_complex_request_exports_long_troubleshooting_and_bus_coverage(tmp_path, monkeypatch):
    root = Path(__file__).resolve().parents[2]
    request = root / "examples/ppt-request-complex.json"
    payload = json.loads(request.read_text(encoding="utf-8"))
    replies = {
        "题目概览": "参数化8位数据选择器\nQuartus II 13.1工程方案\n仅讨论设计与预期验证",
        "背景与目标": "a、b、y均为WIDTH位总线\nsel为一位选择信号\n设计不含时钟与寄存器",
        "步骤一": "建立独立工程目录\n保存Verilog源文件\n器件待实验条件确认",
        "步骤三": "将源文件加入工程\n设置mux2_bus为顶层\n计划执行编译检查",
        "步骤四": "单比特穷举8种\n8位总线共131072种\n逐项比较预期输出\n实际仿真待执行",
        "关键截图与标注": "待观察sel切换\n待核对y与a/b的对应\n同输入时预期输出不变",
        "对比分析": "定向测试检查边界\n穷举覆盖指定参数实例\n参数改变需要独立复测",
        "典型问题与排错": (
            "选择方向反写：核对sel为零时选择a，为一时选择b，避免三目运算分支颠倒\n"
            "总线宽度不匹配：检查a、b和y声明一致，确认WIDTH参数为正整数\n"
            "顶层选择错误：确认工程顶层为mux2_bus，并检查设计文件已加入工程"
        ),
        "讨论焦点": (
            "sel按选择关系切换整条总线\n"
            "WIDTH=1有8种，WIDTH=8有131072种\n"
            "功能验证检查逻辑，时序验证还需分析传播延迟"
        ),
        "总结与延伸": "给出设计与验证方案\n实际编译及仿真待执行\n16位参数需独立复测",
    }

    class LLM:
        def complete_system(self, system, user):
            section = re.search(r"本次只写：\*\*(.*?)\*\*", user)[1]
            return replies[section]

        def complete(self, prompt):
            raise AssertionError("有效复杂输入不应触发重写")

    monkeypatch.setattr(llm_client, "OpenAICompatibleClient", LLM)
    monkeypatch.setattr(Orchestrator, "_output_root", lambda self: None)
    result = cli.main(
        [
            "run",
            "--skill",
            "hnu-digital-circuit-seminar-ppt",
            "--input",
            str(request),
            "--output-dir",
            str(tmp_path),
        ]
    )
    assert result == 0
    prs = Presentation(tmp_path / payload["out_path"])
    assert len(prs.slides) == 12
    text = "\n".join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    assert all(point in text for point in replies["典型问题与排错"].splitlines())
    assert "131072" in text
    assert text.count(payload["collected"]["code_snippet"]) == 1


def test_example_request_cover_never_calls_llm_and_exports_twelve_pages(tmp_path, monkeypatch):
    root = Path(__file__).resolve().parents[2]
    request = root / "examples/ppt-request.json"
    payload = json.loads(request.read_text(encoding="utf-8-sig"))
    replies = {
        "题目概览": "设计二选一选择器",
        "背景与目标": "以选择信号控制输出",
        "步骤一": "建立Quartus工程",
        "步骤二": f"```verilog\n{payload['collected']['code_snippet']}\n```",
        "步骤三": "添加文件并编译检查",
        "步骤四": "遍历三输入的8种组合",
        "关键截图与标注": "观察a、b、sel、y",
        "对比分析": "功能验证核对逻辑\n时序验证分析延迟",
        "典型问题与排错": "核对端口连接和模块名",
        "讨论焦点": "sel为0选a，为1选b",
        "总结与延伸": "预期结果需经实际仿真确认",
    }
    calls = []
    attempts = Counter()
    # Captured failure patterns from the user's generated deck, not idealized prose.
    malformed = {
        "题目概览": """```python
# - 完成实验任务并验证核心逻辑功能
# - 掌握 HDL 设计流程与仿真方法
# - 明确重点讨论问题的深度与排错点
```""",
        "背景与目标": """```python
# - 明确二选一多路选择器逻辑目标
# - 掌握 Quartus II 13.1 环境配置
# - 建立 Verilog 设计到功能验证流程
```""",
        "步骤一": """```json
{
  "title": "工程环境搭建",
  "content": [
    "建立Quartus工程",
    "添加Verilog文件",
    "确认工程目录"
  ],
  "notes": "讲解45秒"
}
```""",
        "步骤三": """```python
# 步骤三：编写与运行功能仿真
# - 编写 Verilog Testbench 并例化模块
# - 构造 a、b、sel 信号的激励组合
# - 预期波形符合多路选择器逻辑
```""",
        "步骤四": """```python
# 步骤四：引脚分配与上板下载
# 检查 Quartus Pin Planner 中的引脚约束关系
# 验证 IO 电平标准与开发板的匹配性
# 预期结果：编译通过且无引脚冲突报错
```""",
    }

    class LLM:
        def complete_system(self, system, user):
            section = re.search(r"本次只写：\*\*(.*?)\*\*", user)[1]
            assert section != "封面", "已采集封面资料不应发送给模型"
            assert section != "步骤二", "已提供的Verilog设计应直接保留，不应让模型重新生成"
            calls.append(section)
            attempts[section] += 1
            if attempts[section] == 1 and section in malformed:
                return malformed[section]
            return replies[section]

        def complete(self, prompt):
            raise AssertionError("这组有效输入不应触发重写")

    monkeypatch.setattr(llm_client, "OpenAICompatibleClient", LLM)
    monkeypatch.setattr(Orchestrator, "_output_root", lambda self: None)
    destination = tmp_path / "outputs"
    result = cli.main(
        [
            "run",
            "--skill",
            "hnu-digital-circuit-seminar-ppt",
            "--input",
            str(request),
            "--output-dir",
            str(destination),
        ]
    )
    assert result == 0
    prs = Presentation(destination / "mux2-seminar.pptx")
    assert len(prs.slides) == 12
    assert set(calls) == set(replies) - {"步骤二"}
    assert attempts["步骤三"] == 2 and attempts["步骤四"] == 2
    cover = "\n".join(s.text for s in prs.slides[0].shapes if s.has_text_frame)
    assert cover.count(payload["collected"]["topic_title"]) == 1
    assert payload["collected"]["member_info"] in cover
    all_text = "\n".join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    assert "# -" not in all_text and '"content"' not in all_text and "讲解45秒" not in all_text
    assert "slide.shapes" not in all_text and "上板下载" not in all_text
    code = payload["collected"]["code_snippet"]
    assert all_text.count(code) == 1
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip() and shape.shape_id in (4, 9, 14):
                if shape.text != code and slide != prs.slides[0]:
                    assert all(
                        r.font.size.pt >= 22 for p in shape.text_frame.paragraphs for r in p.runs
                    )
