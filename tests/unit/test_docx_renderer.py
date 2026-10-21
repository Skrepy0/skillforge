"""DocxRenderer：功能验证 + 确认不再执行 skills 目录代码。"""

from __future__ import annotations

import ast
from pathlib import Path

import pytest

pytest.importorskip("docx")

from docx import Document  # noqa: E402
from skillforge.l2_capability.renderer.docx_renderer import DocxRenderer  # noqa: E402
from skillforge.shared.errors import RenderError  # noqa: E402
from skillforge.shared.types import DraftDocument  # noqa: E402


def _draft() -> DraftDocument:
    return DraftDocument(
        "测试标题",
        [
            {"heading": "第一章", "body": "第一段内容。\n第二段内容。"},
            {"heading": "第二章", "body": "另一段内容。"},
        ],
    )


def _imported_modules(src: str) -> set[str]:
    """提取源码中所有 import 的模块名；注释与 docstring 不计入。"""
    names: set[str] = set()
    for node in ast.walk(ast.parse(src)):
        if isinstance(node, ast.Import):
            for alias in node.names:
                names.add(alias.name)
        elif isinstance(node, ast.ImportFrom) and node.module:
            names.add(node.module)
    return names


def test_render_produces_valid_docx(tmp_path: Path):
    out = DocxRenderer().render(_draft(), tmp_path / "out.docx")
    assert out.exists()
    assert out.stat().st_size > 0

    doc = Document(str(out))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    assert "测试标题" in paragraphs
    assert "第一章" in paragraphs
    assert "第一段内容。" in paragraphs
    assert "另一段内容。" in paragraphs


def test_render_does_not_execute_skills_dir(tmp_path: Path, monkeypatch):
    # 在 skills/utils/docx_common.py 位置放一个会留下标记文件的恶意模块
    malicious_root = tmp_path / "skills"
    (malicious_root / "utils").mkdir(parents=True)
    marker = tmp_path / "executed.txt"
    malicious_code = (
        "from pathlib import Path\n"
        f"Path({str(marker)!r}).write_text('pwned', encoding='utf-8')\n"
        "raise RuntimeError('malicious module executed')\n"
    )
    (malicious_root / "utils" / "docx_common.py").write_text(
        malicious_code,
        encoding="utf-8",
    )

    # 即使把 skills_root 指到这里，渲染器也不读它
    monkeypatch.setenv("SKILLFORGE_SKILLS_ROOT", str(malicious_root))

    out = DocxRenderer().render(_draft(), tmp_path / "out.docx")
    assert out.exists()
    assert not marker.exists(), "渲染器执行了 skills/utils/docx_common.py"
    # 能正常打开，证明用的是受信任模块
    Document(str(out))


def test_render_source_has_no_dynamic_import():
    """静态断言：渲染器与受信任模块都不应 import 动态执行相关模块。"""
    from skillforge.l2_capability.renderer import docx_common, docx_renderer

    for mod in (docx_renderer, docx_common):
        mod_file = mod.__file__
        assert mod_file is not None, f"{mod.__name__} 没有 __file__"
        imported = _imported_modules(Path(mod_file).read_text(encoding="utf-8"))
        for name in imported:
            assert not name.startswith("importlib"), (
                f"{mod.__name__} 仍在 import {name}；渲染器不应动态加载任何代码"
            )


def test_render_wraps_errors_as_render_error(tmp_path: Path, monkeypatch):
    """渲染异常统一包装为 RenderError，调用方无需感知底层异常类型。"""
    from skillforge.l2_capability.renderer import docx_renderer

    def _boom(*_args, **_kwargs):
        raise RuntimeError("boom")

    monkeypatch.setattr(docx_renderer.docx_common, "new_document", _boom)
    with pytest.raises(RenderError, match="docx 渲染失败"):
        DocxRenderer().render(_draft(), tmp_path / "out.docx")
