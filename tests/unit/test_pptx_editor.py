from __future__ import annotations

import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from skillforge.l2_capability.renderer.pptx_editor import PptxEditor
from skillforge.l3_orchestration.pptx_agent import PptxAgent
from skillforge.shared.errors import RenderError


@pytest.fixture
def source(tmp_path):
    path = tmp_path / "template.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Pt(40), Pt(40), Pt(400), Pt(50))
    shape.text = "Template title"
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    prs.save(path)
    return path


def ids(editor):
    slide = editor.snapshot()["slides"][0]
    return {"slide": slide["id"], "shape": slide["shapes"][0]["id"]}


def test_edit_text_style_geometry_and_reopen(source, tmp_path):
    editor = PptxEditor(source)
    address = ids(editor)
    editor.apply(
        [
            {"op": "text", **address, "text": "选择器设计"},
            {
                "op": "style",
                **address,
                "size": 28,
                "font": "微软雅黑",
                "color": "AA3322",
                "bold": True,
                "align": "center",
            },
            {"op": "move", **address, "w": 450, "h": 65},
        ]
    )
    reopened = PptxEditor(editor.save(tmp_path / "edited.pptx"))
    shape = reopened.prs.slides[0].shapes[0]
    run = shape.text_frame.paragraphs[0].runs[0]
    assert shape.text == "选择器设计"
    assert shape.width.pt == 450
    assert run.font.size.pt == 28 and run.font.name == "微软雅黑" and run.font.bold
    assert run._r.rPr.find("a:ea", run._r.nsmap).get("typeface") == "微软雅黑"
    assert run._r.rPr.find("a:ea", run._r.nsmap).get("typeface") == "微软雅黑"
    assert str(run.font.color.rgb) == "AA3322"
    assert Presentation(source).slides[0].shapes[0].text == "Template title"


def test_atomic_batch_does_not_leave_partial_edits(source):
    editor = PptxEditor(source)
    before = editor.snapshot()
    with pytest.raises(RenderError):
        editor.apply(
            [
                {"op": "text", **ids(editor), "text": "changed"},
                {"op": "style", **ids(editor), "size": -1},
            ]
        )
    assert editor.snapshot() == before


@pytest.mark.parametrize(
    "operation",
    [
        {"op": "exec", "code": "print('bad')"},
        {"op": "move", "w": float("nan")},
        {"op": "style", "bold": "false"},
        {"op": "style", "color": "#123456"},
        {"op": "add_image", "asset": "../../secret.png", "x": 1, "y": 1, "w": 20, "h": 20},
        {"op": "delete_slide"},
    ],
)
def test_rejects_invalid_or_unsupported_operations(source, operation):
    editor = PptxEditor(source)
    target = {"slide": ids(editor)["slide"]}
    if operation["op"] in ("move", "style"):
        target["shape"] = ids(editor)["shape"]
    with pytest.raises(RenderError):
        editor.apply([{**target, **operation}])


def test_clone_keeps_stable_ids_and_background(source):
    editor = PptxEditor(source)
    address = ids(editor)
    result = editor.apply([{"op": "clone_slide", "slide": address["slide"]}])
    new_id = result[0]["slide"]
    editor.apply(
        [
            {"op": "text", "slide": new_id, "shape": address["shape"], "text": "Clone"},
            {"op": "delete_slide", "slide": address["slide"]},
        ]
    )
    assert editor.snapshot()["slides"][0]["id"] == new_id
    assert editor.prs.slides[0].shapes[0].text == "Clone"


def test_add_image_shape_text_and_animation(source, tmp_path):
    image = tmp_path / "waveform.png"
    Image.new("RGB", (20, 10), "green").save(image)
    editor = PptxEditor(source, {"waveform": image})
    slide = ids(editor)["slide"]
    result = editor.apply(
        [
            {
                "op": "add_shape",
                "slide": slide,
                "kind": "rect",
                "x": 20,
                "y": 120,
                "w": 100,
                "h": 50,
                "fill": "112233",
            },
            {
                "op": "add_image",
                "slide": slide,
                "asset": "waveform",
                "x": 150,
                "y": 120,
                "w": 100,
                "h": 50,
            },
            {
                "op": "add_text",
                "slide": slide,
                "text": "测试",
                "x": 20,
                "y": 200,
                "w": 200,
                "h": 50,
                "size": 20,
            },
            {"op": "fade_in", **ids(editor), "duration_ms": 700, "delay_ms": 200},
        ]
    )
    prs = Presentation(editor.save(tmp_path / "animated.pptx"))
    assert len(prs.slides[0].shapes) == 4
    assert prs.slides[0].shapes[2].image.blob == image.read_bytes()
    timing = prs.slides[0]._element.find("p:timing", prs.slides[0]._element.nsmap)
    assert timing is not None
    assert timing.xpath(".//p:animEffect")[0].get("filter") == "fade"
    assert timing.xpath(".//p:spTgt")[0].get("spid") == str(ids(editor)["shape"])
    assert len(result) == 4


def test_inspection_flags_geometry_and_text_overflow(source):
    editor = PptxEditor(source)
    editor.apply(
        [
            {"op": "text", **ids(editor), "text": "超长标题" * 30},
            {"op": "move", **ids(editor), "x": 2000},
        ]
    )
    kinds = {issue["kind"] for issue in editor.snapshot()["issues"]}
    assert kinds == {"possible_text_overflow", "outside_canvas"}


class ScriptedLLM:
    def __init__(self, replies):
        self.replies = iter(replies)
        self.prompts = []

    def complete_system(self, system, user):
        self.prompts.append(json.loads(user))
        return json.dumps(next(self.replies), ensure_ascii=False)

    def complete(self, prompt):
        raise AssertionError("system and user should be separate")


def test_agent_reads_edits_then_receives_updated_state(source, tmp_path):
    address = ids(PptxEditor(source))
    llm = ScriptedLLM(
        [
            {
                "operations": [{"op": "text", **address, "text": "Short new title"}],
                "done": False,
                "summary": "改标题",
            },
            {"operations": [], "done": True, "summary": "已检查"},
        ]
    )
    result = PptxAgent(llm).run(source, tmp_path / "out.pptx", {"request": "改标题"})
    assert result["status"] == "delivered" and result["rounds"] == 2
    assert llm.prompts[0]["snapshot"]["slides"][0]["shapes"][0]["text"] == "Template title"
    assert llm.prompts[1]["snapshot"]["slides"][0]["shapes"][0]["text"] == "Short new title"
    assert llm.prompts[1]["feedback"]["results"] == [address]
    assert Path(result["audit_path"]).is_file()
    assert result["visual_review"] is False


def test_agent_recovers_from_bad_operation(source, tmp_path):
    llm = ScriptedLLM(
        [
            {
                "operations": [{"op": "move", **ids(PptxEditor(source)), "w": -5}],
                "done": False,
                "summary": "invalid",
            },
            {"operations": [], "done": True, "summary": "checked original"},
        ]
    )
    result = PptxAgent(llm).run(source, tmp_path / "out.pptx", {})
    assert result["status"] == "delivered"
    assert "error" in llm.prompts[1]["feedback"]


def test_agent_cannot_deliver_uninspected_edits_or_overflow(source, tmp_path):
    address = ids(PptxEditor(source))
    llm = ScriptedLLM(
        [
            {
                "operations": [{"op": "text", **address, "text": "超长" * 200}],
                "done": True,
                "summary": "premature",
            },
            {
                "operations": [{"op": "text", **address, "text": "超长" * 200}],
                "done": False,
                "summary": "overflow",
            },
            {"operations": [], "done": True, "summary": "incorrect"},
        ]
    )
    target = tmp_path / "out.pptx"
    result = PptxAgent(llm, max_rounds=3).run(source, target, {})
    assert result["status"] == "needs_review"
    assert not target.exists()
    assert Path(result["output_path"]).name == "out.review.pptx"
    assert result["issues"]


def test_preview_fails_explicitly_when_renderer_missing(monkeypatch):
    monkeypatch.setattr("skillforge.l3_orchestration.pptx_agent.shutil.which", lambda name: None)
    with pytest.raises(RenderError, match="LibreOffice"):
        PptxAgent(ScriptedLLM([]), preview=True)


def test_agent_never_overwrites_source(source):
    with pytest.raises(RenderError, match="另存"):
        PptxAgent(ScriptedLLM([])).run(source, source, {})


def test_agent_rejects_wrong_generation_page_count(source, tmp_path):
    plan = {"operations": [], "done": True, "summary": "incorrect completion"}
    llm = ScriptedLLM([plan, plan])
    result = PptxAgent(llm, max_rounds=2).run(
        source, tmp_path / "out.pptx", {"sections": ["封面", "原理"]}
    )
    assert result["status"] == "needs_review"
    assert "页数不符合预算" in llm.prompts[1]["feedback"]["error"]


def test_agent_accepts_whole_fenced_json_from_compatible_endpoint(source, tmp_path):
    class FencedLLM:
        def complete(self, prompt):
            return '```json\n{"operations":[],"done":true,"summary":"checked"}\n```'

    result = PptxAgent(FencedLLM()).run(source, tmp_path / "out.pptx", {})
    assert result["status"] == "delivered"


def test_visual_agent_renders_changed_page_again_before_completion(source, tmp_path, monkeypatch):
    rendered = []

    class Preview:
        def render(self, editor, folder, slides):
            rendered.append((slides, editor.prs.slides[0].shapes[0].text))
            path = folder / "page.png"
            Image.new("RGB", (10, 10), "white").save(path)
            return [path]

    class VisualLLM(ScriptedLLM):
        def complete_visual(self, system, user, images):
            assert images[0].read_bytes().startswith(b"\x89PNG")
            return self.complete_system(system, user)

    monkeypatch.setattr("skillforge.l3_orchestration.pptx_agent.SlidePreview", Preview)
    llm = VisualLLM(
        [
            {
                "operations": [{"op": "text", **ids(PptxEditor(source)), "text": "Changed"}],
                "done": False,
                "summary": "edit",
            },
            {"operations": [], "done": True, "summary": "checked preview"},
        ]
    )
    result = PptxAgent(llm, preview=True).run(source, tmp_path / "out.pptx", {})
    assert result["visual_review"] is True
    assert [text for _, text in rendered] == ["Template title", "Changed"]


def test_orchestrator_agent_uses_template_instead_of_prose_generator(source, tmp_path, monkeypatch):
    from skillforge.l2_capability.loader import LoadedSkill
    from skillforge.l3_orchestration.orchestrator import Orchestrator
    from skillforge.shared.types import SessionPhase, SkillContract

    llm = ScriptedLLM([{"operations": [], "done": True, "summary": "keep template"}])
    orch = Orchestrator.from_config(llm=llm)
    contract = SkillContract("ppt-test", "", ["topic"], "pptx", ["封面"], template="x.pptx")
    monkeypatch.setattr(orch.loader, "load", lambda name: LoadedSkill(contract, "参考数据", {}))
    monkeypatch.setattr(orch, "_resolve_template", lambda *args: source)
    monkeypatch.setattr(orch, "_output_root", lambda: None)
    monkeypatch.setattr(orch.queue, "generate_all", lambda *args: pytest.fail("blind generation"))
    result = orch.run(
        request="保留模板",
        collected={"topic": "test"},
        out_path=tmp_path / "out.pptx",
        pptx_agent=True,
        skill_name="ppt-test",
    )
    assert result["status"] == "delivered" and result["skill"] == "ppt-test"
    assert orch.sessions.store.get(result["session_id"]).phase == SessionPhase.DELIVERED
    assert llm.prompts[0]["task"]["collected"]["topic"] == "test"
