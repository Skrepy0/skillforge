"""Inspectable PPT editing operations. Model output is data, never Python or shell code."""

from __future__ import annotations

import math
import re
import unicodedata
from copy import deepcopy
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from skillforge.l2_capability.renderer.pptx_renderer import _clone_slide, _replace_styled
from skillforge.shared.errors import RenderError


def _number(value, name, low=0, high=10000):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise RenderError(f"{name} 必须是数值")
    if not math.isfinite(value) or not low <= value <= high:
        raise RenderError(f"{name} 超出允许范围 [{low}, {high}]")
    return value


def _integer(value, name):
    if type(value) is not int or value < 1:
        raise RenderError(f"{name} 必须是正整数")
    return value


def _color(value):
    if not isinstance(value, str) or not re.fullmatch(r"[0-9a-fA-F]{6}", value):
        raise RenderError("颜色必须是6位十六进制字符串")
    return RGBColor.from_string(value.upper())


def _element(tag, parent=None, **attributes):
    node = OxmlElement(f"p:{tag}")
    for key, value in attributes.items():
        node.set(key, str(value))
    if parent is not None:
        parent.append(node)
    return node


class PptxEditor:
    """Stable slide/shape identifiers, atomic batches, and explicit asset capabilities.

    Geometry is in points (72 points per inch). Slide IDs survive deletion/reordering.
    Local images are supplied by the caller; the model can only select their aliases.
    """

    OPERATIONS = {
        "text": {"slide", "shape", "text"},
        "style": {"slide", "shape", "size", "font", "bold", "color", "align"},
        "move": {"slide", "shape", "x", "y", "w", "h"},
        "add_text": {"slide", "text", "x", "y", "w", "h", "size", "font", "color"},
        "add_shape": {"slide", "kind", "x", "y", "w", "h", "fill", "line"},
        "add_image": {"slide", "asset", "x", "y", "w", "h"},
        "delete_shape": {"slide", "shape"},
        "clone_slide": {"slide"},
        "delete_slide": {"slide"},
        "fade_in": {"slide", "shape", "duration_ms", "delay_ms"},
    }

    def __init__(self, path: Path, assets: dict[str, Path] | None = None):
        self.prs = Presentation(str(path))
        self.assets = {}
        for alias, asset in (assets or {}).items():
            if not isinstance(alias, str) or not isinstance(asset, Path) or not asset.is_file():
                raise RenderError("图片素材必须是已存在的本地文件并有字符串别名")
            if asset.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
                raise RenderError("图片素材仅支持 PNG/JPEG")
            if asset.stat().st_size > 20 * 1024 * 1024:
                raise RenderError("单张素材不能超过20MB")
            self.assets[alias] = asset.read_bytes()

    def _slide(self, slide_id):
        _integer(slide_id, "slide")
        for slide in self.prs.slides:
            if slide.slide_id == slide_id:
                return slide
        raise RenderError(f"幻灯片不存在: {slide_id}")

    def _shape(self, slide, shape_id):
        _integer(shape_id, "shape")
        for shape in slide.shapes:
            if shape.shape_id == shape_id:
                return shape
        raise RenderError(f"元素不存在: {shape_id}")

    def snapshot(self):
        slides, issues = [], []
        width, height = self.prs.slide_width.pt, self.prs.slide_height.pt
        for slide in self.prs.slides:
            shapes = []
            for shape in slide.shapes:
                geometry = dict(
                    zip(
                        ("x", "y", "w", "h"),
                        (
                            round(v.pt, 2)
                            for v in (shape.left, shape.top, shape.width, shape.height)
                        ),
                    )
                )
                item = {
                    "id": shape.shape_id,
                    "name": shape.name,
                    "type": str(shape.shape_type),
                    **geometry,
                }
                if shape.has_text_frame:
                    tf = shape.text_frame
                    item.update(
                        text=shape.text,
                        wrap=tf.word_wrap,
                        margins=[
                            v.pt
                            for v in (
                                tf.margin_left,
                                tf.margin_top,
                                tf.margin_right,
                                tf.margin_bottom,
                            )
                        ],
                        runs=[
                            {
                                "text": r.text,
                                "font": r.font.name,
                                "size": r.font.size.pt if r.font.size else None,
                                "bold": r.font.bold,
                            }
                            for p in tf.paragraphs
                            for r in p.runs
                        ],
                    )
                    if self._text_overflow(shape):
                        issues.append(
                            {
                                "slide": slide.slide_id,
                                "shape": shape.shape_id,
                                "kind": "possible_text_overflow",
                            }
                        )
                if (
                    shape.left < 0
                    or shape.top < 0
                    or shape.left + shape.width > self.prs.slide_width
                    or shape.top + shape.height > self.prs.slide_height
                ):
                    issues.append(
                        {"slide": slide.slide_id, "shape": shape.shape_id, "kind": "outside_canvas"}
                    )
                shapes.append(item)
            slides.append(
                {"id": slide.slide_id, "layout": slide.slide_layout.name, "shapes": shapes}
            )
        return {
            "units": "pt",
            "width": width,
            "height": height,
            "slides": slides,
            "assets": list(self.assets),
            "issues": issues,
            "inspection": "结构检查；字体继承或替换可能影响实际渲染，溢出检测是估算",
        }

    @staticmethod
    def _text_overflow(shape):
        tf = shape.text_frame
        width = max(1, (shape.width - tf.margin_left - tf.margin_right) / 12700)
        height = max(1, (shape.height - tf.margin_top - tf.margin_bottom) / 12700)
        total = 0
        for p in tf.paragraphs:
            if not p.text:
                continue
            # Unknown/inherited metrics cannot be conclusively checked here.
            sizes = [r.font.size.pt for r in p.runs if r.font.size]
            if not sizes:
                continue
            size = max(sizes)
            length = (
                sum(1.05 if unicodedata.east_asian_width(c) in "WF" else 0.65 for c in p.text)
                * size
            )
            if tf.word_wrap is False and length > width:
                return True
            lines = math.ceil(length / width) if tf.word_wrap is not False else 1
            total += max(1, lines) * size * 1.2
        return total > height

    def apply(self, operations: list[dict]):
        if not isinstance(operations, list) or len(operations) > 100:
            raise RenderError("每批 operations 必须为最多100条操作的数组")
        buffer = BytesIO()
        self.prs.save(buffer)
        candidate = object.__new__(PptxEditor)
        buffer.seek(0)
        candidate.prs = Presentation(buffer)
        candidate.assets = self.assets
        results = []
        try:
            for operation in operations:
                results.append(candidate._apply_one(operation))
        except (ValueError, TypeError, KeyError, AttributeError, OverflowError, OSError) as exc:
            raise RenderError("编辑参数非法；本批操作已全部撤销") from exc
        self.prs = candidate.prs
        return results

    def _apply_one(self, op):
        if not isinstance(op, dict) or op.get("op") not in self.OPERATIONS:
            raise RenderError("未知 PPT 编辑操作")
        kind = op["op"]
        if set(op) - self.OPERATIONS[kind] - {"op"}:
            raise RenderError(f"{kind} 含有未知参数")
        slide = self._slide(op["slide"])
        if kind == "clone_slide":
            if len(self.prs.slides) >= 120:
                raise RenderError("最多120页")
            new = _clone_slide(self.prs, slide)
            # Preserve existing animations when cloning; shape identifiers are retained.
            for tag in ("transition", "timing"):
                for element in slide._element.findall(f"{{{slide._element.nsmap['p']}}}{tag}"):
                    new._element.insert_element_before(deepcopy(element), "p:extLst")
            return {"slide": new.slide_id}
        if kind == "delete_slide":
            if len(self.prs.slides) <= 1:
                raise RenderError("不能删除最后一页")
            for slide_id in list(self.prs.slides._sldIdLst):
                if int(slide_id.id) == op["slide"]:
                    self.prs.part.drop_rel(slide_id.rId)
                    self.prs.slides._sldIdLst.remove(slide_id)
                    break
            return {"deleted_slide": op["slide"]}
        shape = None
        if kind.startswith("add_"):
            x, y, w, h = self._box(op)
            if kind == "add_text":
                shape = slide.shapes.add_textbox(x, y, w, h)
                self._text(shape, op["text"])
                self._style(shape, {k: op[k] for k in ("size", "font", "color") if k in op})
            elif kind == "add_shape":
                types = {
                    "rect": MSO_SHAPE.RECTANGLE,
                    "round_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
                    "ellipse": MSO_SHAPE.OVAL,
                    "arrow": MSO_SHAPE.RIGHT_ARROW,
                }
                if op.get("kind") not in types:
                    raise RenderError("形状仅支持 rect/round_rect/ellipse/arrow")
                shape = slide.shapes.add_shape(types[op["kind"]], x, y, w, h)
                if "fill" in op:
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = _color(op["fill"])
                if "line" in op:
                    shape.line.color.rgb = _color(op["line"])
            else:
                if op.get("asset") not in self.assets:
                    raise RenderError("图片必须引用 assets 中的已提供素材别名")
                shape = slide.shapes.add_picture(BytesIO(self.assets[op["asset"]]), x, y, w, h)
        else:
            shape = self._shape(slide, op["shape"])
            if kind == "text":
                self._text(shape, op["text"])
            elif kind == "style":
                self._style(shape, op)
            elif kind == "move":
                for key, attr in (("x", "left"), ("y", "top"), ("w", "width"), ("h", "height")):
                    if key in op:
                        value = _number(op[key], key, 1 if key in ("w", "h") else 0)
                        setattr(shape, attr, Pt(value))
            elif kind == "delete_shape":
                if any(
                    node.get("spid") == str(shape.shape_id)
                    for node in slide._element.iter()
                    if node.tag.endswith("}spTgt")
                ):
                    raise RenderError("该元素被动画引用，暂不支持删除")
                slide.shapes._spTree.remove(shape.element)
            elif kind == "fade_in":
                self._fade(slide, shape, op)
        return {"slide": slide.slide_id, "shape": shape.shape_id}

    @staticmethod
    def _box(op):
        return tuple(
            Pt(_number(op[k], k, 1 if k in ("w", "h") else 0)) for k in ("x", "y", "w", "h")
        )

    @staticmethod
    def _text(shape, text):
        if not shape.has_text_frame or not isinstance(text, str) or len(text) > 10000:
            raise RenderError("文字必须是最多10000字符的字符串，目标必须是文本元素")
        _replace_styled(shape, text)
        shape.text_frame.auto_size = MSO_AUTO_SIZE.NONE

    @staticmethod
    def _style(shape, op):
        if not shape.has_text_frame:
            raise RenderError("该元素不支持文字样式")
        if "size" in op:
            _number(op["size"], "size", 8, 96)
        if "font" in op and (not isinstance(op["font"], str) or not 1 <= len(op["font"]) <= 100):
            raise RenderError("字体名必须是1至100字符的字符串")
        if "bold" in op and type(op["bold"]) is not bool:
            raise RenderError("bold 必须是布尔值")
        aligns = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
        if "align" in op and op["align"] not in aligns:
            raise RenderError("align 必须是 left/center/right")
        color = _color(op["color"]) if "color" in op else None
        for paragraph in shape.text_frame.paragraphs:
            if "align" in op:
                paragraph.alignment = aligns[op["align"]]
            for run in paragraph.runs:
                if "size" in op:
                    run.font.size = Pt(op["size"])
                if "font" in op:
                    run.font.name = op["font"]
                    # python-pptx's name setter only writes a:latin. Chinese text
                    # may still use an explicit a:ea font from the template.
                    properties = run._r.get_or_add_rPr()
                    for tag in ("a:ea", "a:cs"):
                        node = properties.find(tag, properties.nsmap)
                        if node is None:
                            node = OxmlElement(tag)
                            properties.insert_element_before(
                                node,
                                *(["a:cs"] if tag == "a:ea" else []),
                                "a:sym",
                                "a:hlinkClick",
                                "a:hlinkMouseOver",
                                "a:extLst",
                            )
                        node.set("typeface", op["font"])
                if "bold" in op:
                    run.font.bold = op["bold"]
                if color is not None:
                    run.font.color.rgb = color

    @staticmethod
    def _fade(slide, shape, op):
        duration = int(_number(op.get("duration_ms", 500), "duration_ms", 100, 10000))
        delay = int(_number(op.get("delay_ms", 0), "delay_ms", 0, 60000))
        timing = slide._element.find("p:timing", slide._element.nsmap)
        if timing is None:
            timing = _element("timing")
            slide._element.insert_element_before(timing, "p:extLst")
            nodes = _element("tnLst", timing)
            root = _element(
                "cTn",
                _element("par", nodes),
                id=1,
                dur="indefinite",
                restart="never",
                nodeType="tmRoot",
            )
            _element("childTnLst", root)
        root = timing.find(".//p:cTn[@nodeType='tmRoot']", timing.nsmap)
        if root is None:
            raise RenderError("现有动画结构不支持追加淡入")
        children = root.find("p:childTnLst", root.nsmap)
        if children is None:
            children = _element("childTnLst", root)
        ids = [int(node.get("id")) for node in timing.iter() if node.tag.endswith("}cTn")]
        next_id = max(ids, default=0) + 1
        effect_node = _element(
            "cTn",
            _element("par", children),
            id=next_id,
            fill="hold",
            presetID=10,
            presetClass="entr",
            presetSubtype=0,
            nodeType="withEffect",
        )
        _element("cond", _element("stCondLst", effect_node), delay=delay)
        effects = _element("childTnLst", effect_node)
        # Set initial visibility as well as the entrance filter for slideshow playback.
        visibility = _element("set", effects)
        behavior = _element("cBhvr", visibility)
        _element("cTn", behavior, id=next_id + 1, dur=1, fill="hold")
        _element("spTgt", _element("tgtEl", behavior), spid=shape.shape_id)
        attr = _element("attrName", _element("attrNameLst", behavior))
        attr.text = "style.visibility"
        _element("strVal", _element("to", visibility), val="visible")
        effect = _element("animEffect", effects, transition="in", filter="fade")
        behavior = _element("cBhvr", effect)
        _element("cTn", behavior, id=next_id + 2, dur=duration, fill="hold")
        _element("spTgt", _element("tgtEl", behavior), spid=shape.shape_id)

    def save(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path
