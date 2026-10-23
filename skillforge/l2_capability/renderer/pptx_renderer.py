"""DraftDocument -> .pptx，支持技能包自带模板。

- template 非 None 时以它为底版，否则用 python-pptx 默认模板
- 布局优先按 pptx_layouts 声明的 name 查找，找不到回退索引
- 标题框按 placeholder type（TITLE/CENTER_TITLE）查找，内容框按 BODY/OBJECT 查找
- 异常统一包装为 RenderError
"""

from __future__ import annotations

import json
import math
import re
import unicodedata
from copy import deepcopy
from pathlib import Path
from typing import TYPE_CHECKING, Any, cast

from skillforge.l2_capability.pptx_content import (
    normalize_slide_content,
    slide_blocks,
    slide_content_errors,
)
from skillforge.shared.errors import RenderError
from skillforge.shared.types import DraftDocument

if TYPE_CHECKING:
    from pptx.shapes.base import BaseShape


def _split_lines(body: str) -> list[str]:
    lines = []
    in_code = False
    for raw in body.splitlines():
        if raw.strip().startswith("```"):
            in_code = not in_code
            continue
        line = raw.rstrip() if in_code else raw.strip()
        if not in_code:
            line = re.sub(r"^#{1,6}\s+|^[-*+]\s+", "", line)
            line = re.sub(r"\*\*(.+?)\*\*|`([^`]+)`", lambda m: m[1] or m[2], line)
        if line.strip():
            lines.append(line)
    return lines


def _pages(lines, shape):
    # Conservative CJK-aware width and explicit line spacing; never truncate content.
    tf = shape.text_frame
    width = max(20, (shape.width - tf.margin_left - tf.margin_right) / 12700)
    height = max(28, (shape.height - tf.margin_top - tf.margin_bottom) / 12700)
    columns = max(1, int(width / 21))
    rows = max(1, int(height / 28))
    page = []
    for line in lines:
        chunk, used = "", 0.0
        for char in line:
            size = 1 if unicodedata.east_asian_width(char) in "WF" else 0.6
            if chunk and used + size > columns:
                page.append(chunk)
                chunk, used = "", 0.0
                if len(page) == rows:
                    yield page
                    page = []
            chunk += char
            used += size
        if chunk:
            page.append(chunk)
        if len(page) == rows:
            yield page
            page = []
    if page:
        yield page


def _set_text(shape: BaseShape | None, text: str) -> bool:
    if shape is None or not shape.has_text_frame:
        return False
    cast(Any, shape).text_frame.text = text
    _fit_title(shape)
    return True


def _resolve_layout(prs, spec, default_idx: int):
    """按 name / int / 数字字符串查找布局，找不到回退 default_idx。"""
    from pptx.slide import SlideLayouts

    layouts: SlideLayouts = prs.slide_layouts
    if spec is None:
        return layouts[default_idx]
    if isinstance(spec, int):
        try:
            return layouts[spec]
        except IndexError:
            return layouts[default_idx]

    name = str(spec).strip()
    for layout in layouts:
        if layout.name == name:
            return layout
    try:
        return layouts[int(name)]
    except (ValueError, IndexError):
        return layouts[default_idx]


def _find_title(slide):
    from pptx.enum.shapes import PP_PLACEHOLDER

    for ph in slide.placeholders:
        if ph.placeholder_format.type in (
            PP_PLACEHOLDER.TITLE,
            PP_PLACEHOLDER.CENTER_TITLE,
        ):
            return ph
    return None


def _find_body(slide):
    from pptx.enum.shapes import PP_PLACEHOLDER

    for ph in slide.placeholders:
        t = ph.placeholder_format.type
        if t in (
            PP_PLACEHOLDER.BODY,
            PP_PLACEHOLDER.OBJECT,
            PP_PLACEHOLDER.SUBTITLE,
        ):
            return ph
    # 回退：非标题的第一个有文本框的占位符
    title = _find_title(slide)
    for ph in slide.placeholders:
        if ph is not title and ph.has_text_frame:
            return ph
    return None


def _replace_styled(shape, text):
    """Keep template paragraph/run properties rather than resetting the text frame."""
    paragraph = shape.text_frame.paragraphs[0]
    prototype = deepcopy(paragraph._p)
    for child in list(prototype):
        if child.tag.endswith("}r") or child.tag.endswith("}br") or child.tag.endswith("}fld"):
            prototype.remove(child)
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    tf = shape.text_frame
    for child in list(tf._txBody):
        if child.tag.endswith("}p"):
            tf._txBody.remove(child)
    for line in text.split("\n"):
        element = deepcopy(prototype)
        tf._txBody.append(element)
        from pptx.text.text import _Paragraph

        para = _Paragraph(element, tf)
        run = para.add_run()
        if run_properties is not None:
            run._r.insert(0, deepcopy(run_properties))
        run.text = line


def _fit_title(shape, *, max_width=None, max_height=None):
    """Use available layout space and wrapping, never shrink titles below 24pt."""
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Pt

    tf = shape.text_frame
    if max_width is not None:
        shape.width = Pt(max_width)
    width = (shape.width - tf.margin_left - tf.margin_right) / 12700 * 0.95
    available_height = max_height if max_height is not None else shape.height / 12700
    height = available_height - (tf.margin_top + tf.margin_bottom) / 12700
    if width <= 0 or height <= 0:
        raise RenderError("标题文本框没有可用空间")
    sizes = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size]
    original = max(24, max(sizes, default=32))
    chosen = None
    required = 0.0
    for size in range(int(original), 23, -1):
        rows = sum(
            max(
                1,
                math.ceil(
                    sum(1.05 if unicodedata.east_asian_width(c) in "WF" else 0.7 for c in p.text)
                    * size
                    / width
                ),
            )
            for p in tf.paragraphs
        )
        required = rows * size * 1.25
        if required <= height:
            chosen = size
            break
    if chosen is None:
        raise RenderError("标题在24磅下仍超出页面空间，请精简标题或调整布局")
    if max_height is not None:
        shape.height = Pt(
            max(shape.height / 12700, required + (tf.margin_top + tf.margin_bottom) / 12700)
        )
    for para in tf.paragraphs:
        para.space_before = Pt(0)
        para.space_after = Pt(0)
        para.line_spacing = 1.25
        for run in para.runs:
            run.font.size = Pt(chosen)
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE


def _clone_slide(prs, source):
    from pptx.oxml.ns import qn

    slide = prs.slides.add_slide(source.slide_layout)
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape.element)
    rels = {}
    for rel in source.part.rels.values():
        if rel.reltype.endswith(("/slideLayout", "/notesSlide")):
            continue
        target = rel.target_ref if rel.is_external else rel.target_part
        rels[rel.rId] = slide.part.relate_to(target, rel.reltype, rel.is_external)
    for shape in source.shapes:
        element = deepcopy(shape.element)
        for node in element.iter():
            for attr in (qn("r:embed"), qn("r:link"), qn("r:id")):
                if node.get(attr) in rels:
                    node.set(attr, rels[node.get(attr)])
        slide.shapes._spTree.insert_element_before(element, "p:extLst")
    if source._element.cSld.bg is not None:
        slide._element.cSld.insert(0, deepcopy(source._element.cSld.bg))
    return slide


def _render_bound(prs, draft, bindings):
    from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Pt

    sources = list(prs.slides)
    original_ids = list(prs.slides._sldIdLst)

    def create(spec, title, cover=False):
        slide = _clone_slide(prs, sources[spec["slide"] - 1])
        shapes = {shape.shape_id: shape for shape in slide.shapes}
        title_shape = shapes[spec["title"]]
        _replace_styled(title_shape, title)
        slots = [shapes[sid] for sid in spec["body"]]
        if cover:
            title_shape.left = Pt(54)
            for para in title_shape.text_frame.paragraphs:
                para.alignment = PP_ALIGN.CENTER
            _fit_title(
                title_shape,
                max_width=prs.slide_width.pt - 108,
                max_height=min(s.top.pt for s in slots) - title_shape.top.pt - 24,
            )
        else:
            _fit_title(
                title_shape,
                max_width=prs.slide_width.pt - title_shape.left.pt - 54,
                max_height=min(s.top.pt for s in slots) - title_shape.top.pt - 20,
            )
        for sid in spec.get("clear", []) + spec["body"]:
            _replace_styled(shapes[sid], "")
        return slots

    def fill(shape, text, size=22, code=False):
        _replace_styled(shape, text)
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        for para in tf.paragraphs:
            para.space_before = Pt(0)
            para.space_after = Pt(0)
            para.line_spacing = 1.25
            for run in para.runs:
                run.font.size = Pt(size)
                if code:
                    run.font.name = "Consolas"

    sections = list(draft.sections)
    cover_text = ""
    if sections and sections[0].get("heading") in ("封面", "Cover"):
        cover = sections.pop(0)
        cover_text = normalize_slide_content(
            cover.get("body", ""), cover["heading"], {"topic_title": draft.title}
        )
        errors = slide_content_errors(cover_text, cover["heading"])
        if errors:
            raise RenderError("封面页面内容不符合预算：" + "；".join(errors))
    cover_slots = create(bindings["cover"], draft.title, cover=True)
    member_shape = cover_slots[0]
    member_shape.left, member_shape.width = Pt(54), Pt(prs.slide_width.pt - 108)
    member_shape.height = Pt(prs.slide_height.pt - member_shape.top.pt - 40)
    if len(cover_text.splitlines()) > 5:
        cover_text = "；".join(cover_text.splitlines())
    tf = member_shape.text_frame
    width = (member_shape.width - tf.margin_left - tf.margin_right) / 12700 * 0.95
    height = (member_shape.height - tf.margin_top - tf.margin_bottom) / 12700
    member_size = None
    for size in (18, 17, 16):
        rows = sum(
            max(
                1,
                math.ceil(
                    sum(1.05 if unicodedata.east_asian_width(c) in "WF" else 0.65 for c in line)
                    * size
                    / width
                ),
            )
            for line in cover_text.splitlines()
        )
        if rows * size * 1.25 <= height:
            member_size = size
            break
    if member_size is None:
        raise RenderError("封面成员资料超出布局容量，请精简资料或使用单独的成员页")
    fill(member_shape, cover_text, size=member_size)
    for paragraph in member_shape.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
    for section in sections:
        heading = section.get("heading", "")
        errors = slide_content_errors(section.get("body", ""), heading)
        if errors:
            raise RenderError(f"页面“{heading}”内容不符合预算：" + "；".join(errors))
        blocks = slide_blocks(section.get("body", ""), heading)
        spec = bindings.get("sections", {}).get(heading, bindings["content"])
        slots = create(spec, heading)
        if len(blocks) > len(slots):
            raise RenderError(f"页面“{heading}”的模板文本槽不足")
        left = min(s.left.pt for s in slots)
        top = min(s.top.pt for s in slots)
        right = max(s.left.pt + s.width.pt for s in slots)
        bottom = max(s.top.pt + s.height.pt for s in slots)
        cursor = top
        columns = spec.get("arrangement") == "columns"
        for index, (slot, block) in enumerate(zip(slots, blocks)):
            size = 16 if block.code else 22
            width = (
                (right - left - 24 * (len(blocks) - 1)) / len(blocks) if columns else right - left
            )
            rows = sum(
                max(
                    1,
                    math.ceil(
                        sum(1.05 if unicodedata.east_asian_width(c) in "WF" else 0.65 for c in line)
                        * size
                        / (width - 20)
                    ),
                )
                for line in block.text.splitlines()
            )
            required = max(50, rows * size * 1.25 + 12)
            if (top if columns else cursor) + required > bottom:
                raise RenderError(f"页面“{heading}”超出模板容量，请精简要点或代码，禁止自动拆页")
            slot.left = Pt(left + index * (width + 24) if columns else left)
            slot.top = Pt(top if columns else cursor)
            slot.width, slot.height = Pt(width), Pt(required)
            fill(slot, block.text, size, block.code)
            cursor += required + 24
    for slide_id in original_ids:
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)


class PptxRenderer:
    def render(
        self,
        draft: DraftDocument,
        out_path: Path,
        template: Path | None = None,
        layouts: dict[str, str | int] | None = None,
    ) -> Path:
        try:
            from pptx import Presentation
            from pptx.util import Pt

            prs = Presentation(str(template)) if template else Presentation()
            bindings_path = Path(template).with_suffix(".bindings.json") if template else None
            if bindings_path and bindings_path.exists():
                bindings = json.loads(bindings_path.read_text(encoding="utf-8-sig"))
                _render_bound(prs, draft, bindings)
                out_path = Path(out_path)
                out_path.parent.mkdir(parents=True, exist_ok=True)
                prs.save(str(out_path))
                return out_path
            # Templates provide masters/layouts, not example content.
            for slide_id in list(prs.slides._sldIdLst):
                prs.part.drop_rel(slide_id.rId)
                prs.slides._sldIdLst.remove(slide_id)
            layout_map = layouts or {}

            title_layout = _resolve_layout(prs, layout_map.get("title"), 0)
            content_layout = _resolve_layout(prs, layout_map.get("content"), 1)

            title_slide = prs.slides.add_slide(title_layout)
            _set_text(_find_title(title_slide), draft.title or "Untitled")
            subtitle = _find_body(title_slide)
            if subtitle is not None and subtitle is not _find_title(title_slide):
                _set_text(subtitle, "Generated by SkillForge")

            for section in draft.sections:
                slide = prs.slides.add_slide(content_layout)
                heading = section.get("heading", "")
                _set_text(_find_title(slide), heading)
                body_shape = _find_body(slide)
                if body_shape is None or not body_shape.has_text_frame:
                    raise RenderError("内容布局缺少正文文本占位符，请选择 Title and Content 布局")
                pages = list(_pages(_split_lines(section.get("body", "")), body_shape))
                for index, lines in enumerate(pages):
                    if index:
                        slide = prs.slides.add_slide(content_layout)
                        _set_text(_find_title(slide), f"{heading}（续 {index + 1}）")
                        body_shape = _find_body(slide)
                    text_frame = body_shape.text_frame
                    text_frame.word_wrap = True
                    text_frame.text = "\n".join(lines)
                    for para in text_frame.paragraphs:
                        para.space_before = Pt(0)
                        para.space_after = Pt(0)
                        para.line_spacing = Pt(28)
                        for run in para.runs:
                            run.font.size = Pt(20)

            out_path = Path(out_path)
            out_path.parent.mkdir(parents=True, exist_ok=True)
            prs.save(str(out_path))
        except Exception as exc:
            raise RenderError(f"pptx 渲染失败: {exc}") from exc
        return out_path
