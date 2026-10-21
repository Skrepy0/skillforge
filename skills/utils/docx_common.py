"""共享的 docx 排版函数，供各技能脚本与渲染器复用。"""

from __future__ import annotations

from typing import TYPE_CHECKING, cast

from docx import Document as create_document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.styles.style import ParagraphStyle

if TYPE_CHECKING:
    from docx.document import Document


def _set_east_asia_font(element, font_name: str) -> None:
    """给 CT_Style 或 CT_R 设置中文（eastAsia）字体，自动补建 rPr/rFonts。"""
    rpr = element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font_name)


def new_document() -> Document:
    doc = create_document()
    style = cast(ParagraphStyle, doc.styles["Normal"])
    style.font.name = "宋体"
    style.font.size = Pt(12)
    _set_east_asia_font(style.element, "宋体")
    return doc


def add_heading(doc: Document, text: str, level: int = 1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "黑体"
        _set_east_asia_font(run._element, "黑体")
    return h


def add_paragraph(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Pt(24)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def save(doc: Document, path: str) -> None:
    doc.save(path)
